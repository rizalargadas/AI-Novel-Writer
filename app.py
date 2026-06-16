from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import hashlib
import re
from pathlib import Path
from dotenv import load_dotenv
import streamlit as st
from openai import OpenAI

# Load API key from .env
load_dotenv()

client = OpenAI()

# Folders
BASE_DIR = Path(".")
OUTPUTS_DIR = BASE_DIR / "outputs"
CHAPTERS_DIR = BASE_DIR / "chapters"
PUBLISHING_DIR = BASE_DIR / "publishing"

OUTPUTS_DIR.mkdir(exist_ok=True)
CHAPTERS_DIR.mkdir(exist_ok=True)
PUBLISHING_DIR.mkdir(exist_ok=True)

WRITING_PHILOSOPHY = """
WRITING PHILOSOPHY (internalize this before every chapter)

These are the craft principles that separate memorable fiction from forgettable prose.
Violating them is the most common reason readers say a story "has no soul."

1. SHOW, DON'T TELL — BUT GO DEEPER THAN THE CLICHÉ
Do not write: "She was nervous."
Write the nervous: the rehearsed sentence dissolving before it reaches her lips,
the way she smooths the same wrinkle in her skirt three times.

Every emotion must have a physical body. Ground feelings in sensation, gesture,
and involuntary reaction.

2. SCENE OVER SUMMARY
Do not summarize what happened. Put the reader inside the moment it happens.
Bad: "They argued for an hour and she left angry."
Good: The argument lives on the page — the specific words, the silences,
the thing she almost said and didn't.

If a moment matters to the character, it must cost the reader something too.
Don't skip the hard scenes. Slow down exactly where it hurts.

3. DIALOGUE MUST DO DOUBLE WORK
Every line of dialogue should carry subtext. Characters rarely say exactly what
they mean — they deflect, overexplain, go quiet at the wrong moment, joke when
they're breaking.

Let silence and indirection do the emotional labor. What is NOT said is often
the most important line.

4. SPECIFICITY IS SOUL
Vague writing feels hollow. Specific writing feels true.
Not: "a nice restaurant" — the exact name, the candle that kept guttering,
the way the menu was printed in a font too small for the lighting.
Not: "she smelled his cologne" — which cologne, what it reminded her of,
whether it made things better or worse.

Sensory and setting-specific detail transforms a plot into a world.
Earn the reader's trust with precision.

5. CHARACTER INTERIORITY — THE INNER LIFE IS THE REAL STORY
Between every action and reaction, there is a thought, a memory, a contradiction.
Let the reader live inside the character's head — not just observing them,
but thinking alongside them.

Readers don't fall in love with plot events. They fall in love with a character's
specific way of seeing the world. Protect and develop that voice relentlessly.

6. PACING IS EMOTIONAL, NOT MECHANICAL
Short sentences accelerate. Long sentences breathe. Paragraph breaks create
weight. A single sentence on its own line can stop a reader's heart.

Vary your rhythm intentionally. Tension scenes move fast. Intimate scenes slow
almost to stillness. Do not write every scene at the same tempo.

7. AVOID THESE SPECIFIC FAILURE MODES
- Emotion-labeling: "She felt sad." / "He was furious." Always replace with
  physical or behavioral rendering.
- Chapter padding: scenes that exist only to relay information, with no emotional stakes.
- On-the-nose dialogue: characters explaining feelings directly in ways real people never do.
- Symmetrical sentence structures that create a robotic rhythm.
- Adjective stacking instead of precise nouns and active verbs.
"""

def split_options(text, label="OPTION"):
    """
    Splits AI output into selectable chunks.
    Example labels: OPTION or ENDING.
    It ignores the RECOMMENDED OPTION / RECOMMENDED ENDING section.
    """
    pattern = rf"(?is)\b{label}\s+\d+\b.*?(?=\n\s*{label}\s+\d+\b|\n\s*RECOMMENDED\s+(OPTION|ENDING)\s*:|$)"
    parts = re.findall(pattern, text)

    # re.findall returns tuples because of the group in the lookahead.
    # This fixes that issue by using re.finditer instead.
    matches = re.finditer(pattern, text)
    clean_parts = [match.group(0).strip() for match in matches]

    return clean_parts

def split_fixed_chapter_and_log(audit_fix_text):
    """
    Splits AI output into fixed chapter and change log.
    Expected sections:
    # FIXED CHAPTER
    # CHANGE LOG
    """
    fixed_marker = "# FIXED CHAPTER"
    log_marker = "# CHANGE LOG"

    if fixed_marker in audit_fix_text and log_marker in audit_fix_text:
        fixed_part = audit_fix_text.split(fixed_marker, 1)[1].split(log_marker, 1)[0].strip()
        log_part = audit_fix_text.split(log_marker, 1)[1].strip()
        return fixed_part, log_part

    # fallback: keep original output as log if split fails
    return audit_fix_text, "Could not split fixed chapter and change log cleanly."

def split_fixed_chapter_and_final_log(audit_fix_text):
    """
    Splits final audit output into fixed chapter and final audit change log.
    Expected sections:
    # FIXED CHAPTER
    # FINAL AUDIT CHANGE LOG
    """
    fixed_marker = "# FIXED CHAPTER"
    log_marker = "# FINAL AUDIT CHANGE LOG"

    if fixed_marker in audit_fix_text and log_marker in audit_fix_text:
        fixed_part = audit_fix_text.split(fixed_marker, 1)[1].split(log_marker, 1)[0].strip()
        log_part = audit_fix_text.split(log_marker, 1)[1].strip()
        return fixed_part, log_part

    return audit_fix_text, "Could not split fixed chapter and final audit change log cleanly."

def extract_recommended_number(text, recommendation_label="OPTION", item_label="OPTION"):
    """
    Finds the recommended option/ending number from AI output.

    Example:
    RECOMMENDED OPTION:
    Option 4 is the strongest...

    RECOMMENDED ENDING:
    Ending 10 is the strongest...
    """
    recommended_section = re.search(
        rf"(?is)RECOMMENDED\s+{recommendation_label}\s*:.*",
        text
    )

    if not recommended_section:
        return None

    match = re.search(
        rf"(?i)\b{item_label}\s+(\d+)\b",
        recommended_section.group(0)
    )

    if match:
        return int(match.group(1))

    return None

def save_markdown(file_path, content):
    """Save text into a markdown file."""
    file_path = Path(file_path)
    file_path.parent.mkdir(exist_ok=True)
    file_path.write_text(content, encoding="utf-8")


def read_markdown(file_path):
    """Read markdown file if it exists."""
    file_path = Path(file_path)
    if file_path.exists():
        return file_path.read_text(encoding="utf-8")
    return ""


def ask_openai(prompt, model="gpt-5.5"):
    """
    Sends prompt to OpenAI.
    Use this only when the user clicks a button.
    """
    response = client.responses.create(
        model=model,
        reasoning={"effort": "low"},
        input=prompt
    )
    return response.output_text


def generate_pitch_prompt(story_idea, genre):
    return f"""
You are a bestselling LGBTQ+ fiction writer with expertise in {genre}.

Using the details below, generate 5 distinct story plot concepts.

Format each option exactly like this:

OPTION 1
Title:
Logline:
Summary:
Central Emotional Conflict:
Commercial Potential:

OPTION 2
Title:
Logline:
Summary:
Central Emotional Conflict:
Commercial Potential:

Continue until OPTION 5.

After all 5 options, add:
RECOMMENDED OPTION:
Briefly explain which option has the strongest commercial potential and why.

USER INPUT:
{story_idea}
"""


def generate_endings_prompt(pitch_text):
    return f"""
Based on the selected story plot below, create 10 different ending ideas.

Format each ending exactly like this:

ENDING 1
Title:
Ending Summary:
Emotional Payoff:
Reader Satisfaction:
Risk:

ENDING 2
Title:
Ending Summary:
Emotional Payoff:
Reader Satisfaction:
Risk:

Continue until ENDING 10.

After all 10 options, add:
RECOMMENDED ENDING:
Recommend the best ending according to:
- uniqueness
- emotional impact
- commercial reader satisfaction

SELECTED PLOT:
{pitch_text}
"""

def generate_expand_ending_prompt(pitch_text, selected_ending):
    return f"""
Based on the selected plot and selected ending below, expand the ending into a clean Markdown reference document.

This expanded ending will become the official confirmed ending for the novel.

Include:
- Ending Title
- Final Ending Summary
- Emotional Payoff
- Character Resolution
- Final Twist or Reveal, if applicable
- Last Scene Description
- Why This Ending Works
- Notes for Foreshadowing Earlier Chapters

Make the ending emotionally moving, unique, commercially satisfying, and structurally useful for later outlining.

SELECTED PLOT:
{pitch_text}

SELECTED ENDING:
{selected_ending}
"""

def generate_metadata_prompt(pitch_text, ending_text):
    return f"""
Using the story plot and confirmed ending below, create a clean Markdown metadata file.

IMPORTANT:
Start the file with this exact field:

Novel Title: [Insert the final commercial novel title]

Then include:

1. Target Audience
Give a one-sentence explanation of the target audience.

2. Short Story Description
Write one short paragraph. Include:
- genre
- setting
- protagonist age bracket
- protagonist short description

PLOT:
{pitch_text}

CONFIRMED ENDING:
{ending_text}
"""


def generate_character_prompt(metadata_text, ending_text):
    return f"""
You are a genre-savvy fiction editor and character-development expert.

I am writing a novel to self-publish on Draft2Digital.

Story summary:
{metadata_text}

Confirmed ending:
{ending_text}

Create a complete character plan that supports the ending emotionally and structurally.

LEAD CHARACTERS — Full Deep-Dive Profile
For each lead character, provide:
- Name & Age
- Role
- Core Identity
- Wound
- Desire
- Need
- Fear
- Lie They Believe
- Truth They'll Learn
- Personality Snapshot
- Physical Description
- Voice & Mannerisms
- Key Relationships
- Arc

SUPPORTING CHARACTERS — Concise Functional Profile
For each supporting character, provide:
- Name, Age, Nationality
- Role
- Core Identity
- Desire / Fear
- Relationship to leads
- Story Function
- Voice in one sentence

Avoid these names and surnames:
Okafor, Marcus, Voss, Osei, Wren, Calloway, Soren, Sable, Chen, Mara, Dale, Vera, Calhoun, Petra, Maram, Priya, Eli, Nora, Cass, Vesper, Rhea

Format as a clean Markdown reference document.
"""

def extract_title_from_metadata(metadata_text, selected_pitch_text=""):
    """
    Extracts the real novel title.
    Priority:
    1. Novel Title from metadata
    2. Title from selected pitch
    3. Other title fields
    4. Fallback
    """

    metadata_patterns = [
        r"(?im)^Novel Title:\s*(.+)$",
        r"(?im)^\*\*Novel Title:\*\*\s*(.+)$",
        r"(?im)^Book Title:\s*(.+)$",
        r"(?im)^\*\*Book Title:\*\*\s*(.+)$",
        r"(?im)^Title:\s*(.+)$",
        r"(?im)^\*\*Title:\*\*\s*(.+)$",
    ]

    for pattern in metadata_patterns:
        match = re.search(pattern, metadata_text)
        if match:
            title = clean_title_text(match.group(1))
            if title and title.lower() not in ["story metadata", "metadata"]:
                return title

    pitch_patterns = [
        r"(?im)^Title:\s*(.+)$",
        r"(?im)^\*\*Title:\*\*\s*(.+)$",
    ]

    for pattern in pitch_patterns:
        match = re.search(pattern, selected_pitch_text)
        if match:
            title = clean_title_text(match.group(1))
            if title:
                return title

    return "Untitled Novel"


def clean_title_text(title):
    """
    Cleans markdown and placeholder text from title.
    """
    title = re.sub(r"[*_#\[\]]", "", title).strip()
    title = title.replace("Insert the final commercial novel title", "").strip()
    title = title.strip(":-— ")
    return title

def safe_folder_name(name):
    """
    Makes a safe Windows folder name from the novel title.
    """
    name = re.sub(r'[<>:"/\\|?*]', "", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name or "Untitled Novel"


def detect_author_name(metadata_text):
    """
    Uses AERESSA for WLW/sapphic/lesbian stories.
    Uses VoidAndVelvet for MLM/gay male stories.
    Defaults to AERESSA.
    """
    text = metadata_text.lower()

    if any(word in text for word in ["mlm", "gay male", "men loving men", "male/male"]):
        return "VoidAndVelvet"

    return "AERESSA"


def add_markdown_chapter_to_doc(doc, chapter_text, chapter_number, chapter_title):
    """
    Adds one chapter to DOCX.
    Always creates the official chapter title as H1:
    Chapter 1 - Title

    It skips all existing chapter heading lines inside the chapter text,
    whether they are H1, H2, or plain text.
    """

    official_title = f"Chapter {chapter_number} - {chapter_title}".strip(" -")
    doc.add_heading(official_title, level=1)

    lines = chapter_text.splitlines()

    for line in lines:
        clean_line = line.strip()

        if not clean_line:
            continue

        # Skip existing chapter title lines like:
        # # Chapter 1 - Title
        # ## Chapter 1 - Title
        # Chapter 1 - Title
        if re.match(r"(?i)^\s*#{0,6}\s*Chapter\s+\d+\b", clean_line):
            continue

        if clean_line.startswith("### "):
            doc.add_heading(clean_line.replace("### ", "", 1).strip(), level=3)
        elif clean_line.startswith("## "):
            doc.add_heading(clean_line.replace("## ", "", 1).strip(), level=2)
        elif clean_line.startswith("# "):
            doc.add_heading(clean_line.replace("# ", "", 1).strip(), level=2)
        else:
            para = doc.add_paragraph(clean_line)
            para.style = doc.styles["Normal"]

def get_chapter_title_from_outline(outline_text, chapter_number):
    """
    Pulls chapter title from Step 5 outline.
    Looks for:
    Chapter 1 - Title
    # Chapter 1 - Title
    ## Chapter 1 - Title
    """

    pattern = rf"(?im)^\s*#{0,6}\s*Chapter\s+{chapter_number}\s*[-:—]\s*(.+)$"
    match = re.search(pattern, outline_text)

    if match:
        title = match.group(1).strip()
        title = re.sub(r"[*_#\[\]]", "", title).strip()

        if title and not re.match(rf"(?i)^chapter\s+{chapter_number}$", title):
            return title

    return "Untitled"

def compile_chapters_to_docx(metadata_text, selected_pitch_text="", outline_text="", output_path=None):
    """
    Compiles all chapter_XX.md files into one formatted DOCX.
    Ensures:
    - Correct novel title
    - 'by Author Name'
    - Every chapter title is H1
    - Chapter number/title comes from Step 5 outline when possible
    """

    doc = Document()

    styles = doc.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "Garamond"
    normal_style.font.size = Pt(12)

    heading1 = styles["Heading 1"]
    heading1.font.name = "Garamond"
    heading1.font.size = Pt(20)
    heading1.font.bold = True

    novel_title = extract_title_from_metadata(metadata_text, selected_pitch_text)
    author_name = detect_author_name(metadata_text)

    # Front page
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(novel_title)
    title_run.bold = True
    title_run.font.name = "Cinzel"
    title_run.font.size = Pt(40)

    author_para = doc.add_paragraph()
    author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = author_para.add_run(f"by {author_name}")
    author_run.font.name = "Garamond"
    author_run.font.size = Pt(18)

    doc.add_page_break()

    chapter_files = sorted(CHAPTERS_DIR.glob("chapter_*.md"))

    chapter_files = [
        file for file in chapter_files
        if not file.name.endswith("_log.md")
    ]

    for index, chapter_file in enumerate(chapter_files):
        chapter_match = re.search(r"chapter_(\d+)\.md", chapter_file.name, re.I)
        chapter_number = int(chapter_match.group(1)) if chapter_match else index + 1

        chapter_title = get_chapter_title_from_outline(outline_text, chapter_number)

        chapter_text = chapter_file.read_text(encoding="utf-8")
        add_markdown_chapter_to_doc(doc, chapter_text, chapter_number, chapter_title)

        if index < len(chapter_files) - 1:
            doc.add_page_break()

    if output_path is None:
        output_path = PUBLISHING_DIR / "final_manuscript.docx"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc.save(output_path)

    return output_path

def save_text_as_docx(text, output_path, title=None):
    """
    Saves markdown-ish text into a simple DOCX.
    """

    doc = Document()

    styles = doc.styles
    normal_style = styles["Normal"]
    normal_style.font.name = "Garamond"
    normal_style.font.size = Pt(12)

    if title:
        heading = doc.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in text.splitlines():
        clean_line = line.strip()

        if not clean_line:
            continue

        if clean_line.startswith("### "):
            doc.add_heading(clean_line.replace("### ", "", 1).strip(), level=3)
        elif clean_line.startswith("## "):
            doc.add_heading(clean_line.replace("## ", "", 1).strip(), level=2)
        elif clean_line.startswith("# "):
            doc.add_heading(clean_line.replace("# ", "", 1).strip(), level=1)
        else:
            doc.add_paragraph(clean_line)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)

    return output_path

def export_final_publishing_outputs(parent_folder, metadata_text, selected_pitch_text, outline_text, cover_prompt, d2d_metadata, youtube_metadata):
    """
    Exports the 4 final publishing files into:
    [chosen parent folder] / [Novel Title] /

    Outputs:
    - final full novel .docx
    - book cover prompt .txt
    - Draft2Digital details .docx
    - YouTube details .docx
    """

    novel_title = extract_title_from_metadata(metadata_text, selected_pitch_text)
    novel_folder_name = safe_folder_name(novel_title)

    export_dir = Path(parent_folder) / novel_folder_name
    export_dir.mkdir(parents=True, exist_ok=True)

    manuscript_path = compile_chapters_to_docx(
        metadata_text=metadata_text,
        selected_pitch_text=selected_pitch_text,
        outline_text=outline_text,
        output_path=export_dir / f"{novel_folder_name} - Final Manuscript.docx"
    )

    cover_path = export_dir / f"{novel_folder_name} - Book Cover Prompt.txt"
    cover_path.write_text(cover_prompt, encoding="utf-8")

    d2d_path = save_text_as_docx(
        d2d_metadata,
        export_dir / f"{novel_folder_name} - Draft2Digital Details.docx",
        title="Draft2Digital Details"
    )

    youtube_path = save_text_as_docx(
        youtube_metadata,
        export_dir / f"{novel_folder_name} - YouTube Details.docx",
        title="YouTube Details"
    )

    return {
        "folder": export_dir,
        "manuscript": manuscript_path,
        "cover_prompt": cover_path,
        "draft2digital": d2d_path,
        "youtube": youtube_path,
    }

def generate_chapter_audit_and_fix_prompt(chapter_number, chapter_text, metadata_text, character_text, ending_text, outline_text):
    return f"""
You are a professional fiction continuity editor and prose revision specialist.

Audit Chapter {chapter_number} against the project files.

If you find errors, contradictions, weak prose, missing emotional stakes, timeline issues, or craft issues, fix them directly.

Your output must have exactly two sections:

# FIXED CHAPTER

[Write the complete corrected chapter here.]

# CHANGE LOG

List every meaningful fix made.

For each fix, include:
- Original problem
- Fix applied
- Why the fix was necessary

Check for:
- Plot continuity issues
- Character consistency issues
- Timeline problems
- Logic problems
- Contradictions with metadata, character profiles, ending, or outline
- Weak prose
- Emotion-labeling instead of showing
- On-the-nose dialogue
- Chapter padding
- Missing emotional stakes
- Robotic rhythm
- Adjective stacking

If no issues are found, keep the chapter unchanged and say in the change log:
No major issues found. Chapter retained as originally drafted.

PROJECT FILES:

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}

CONFIRMED ENDING:
{ending_text}

OUTLINE:
{outline_text}

CHAPTER TEXT:
{chapter_text}
"""

def generate_outline_prompt(metadata_text, character_text, ending_text):
    return f"""
You are a professional story architect and developmental editor.

Based on the story materials below, choose the strongest story structure from:
- Freytag's Pyramid
- The Hero's Journey
- Three Act Structure
- Dan Harmon's Story Circle
- Fichtean Curve
- Save the Cat Beat Sheet
- Seven-Point Story Structure

Then create a chapter-by-chapter outline.

IMPORTANT:
- The outline must have at least 20 chapters.
- If the story needs more than 20 chapters, create more.
- Each chapter heading must start with this exact format:
Chapter 1 - [Chapter Title]
Chapter 2 - [Chapter Title]
Chapter 3 - [Chapter Title]

Each chapter should include:
- Hooky chapter title
- Most important scene
- Main emotion
- Purpose
- Two-paragraph chapter description

Write as a clean Markdown reference document.

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}

CONFIRMED ENDING:
{ending_text}
"""


def generate_chapter_prompt(chapter_number, metadata_text, character_text, ending_text, outline_text):
    return f"""
Write Chapter {chapter_number} of the novel.

PROJECT CONTEXT:
Metadata:
{metadata_text}

Character Profiles:
{character_text}

Confirmed Ending:
{ending_text}

Outline:
{outline_text}

{WRITING_PHILOSOPHY}

PRE-WRITING REVIEW:
Before writing, silently answer:
- What is the emotional core of this chapter?
- What does the reader need to feel by the end?
- What changes emotionally between the beginning and end of this chapter?

WRITING RULES:
- Write only the chapter.
- Format the chapter name as H1, example: # Chapter {chapter_number} - Chapter Title
- Do not summarize feelings; render them in scene.
- Use sensory detail, physical reaction, interiority, subtext, and emotionally specific dialogue.
- Do not include the audit in the chapter file.
- Do not use emojis.
- Do not use horizontal lines.
- Maintain continuity with the project documents.
- Mature scenes are permitted, but do not write erotica.
"""

def generate_chapter_audit_prompt(chapter_number, chapter_text, metadata_text, character_text, ending_text, outline_text):
    return f"""
You are a professional fiction continuity editor and prose quality auditor.

Audit Chapter {chapter_number} against the project files.

Check for:
- Plot continuity issues
- Character consistency issues
- Timeline problems
- Logic problems
- Contradictions with metadata, character profiles, ending, or outline
- Weak prose
- Emotion-labeling instead of showing
- On-the-nose dialogue
- Chapter padding
- Missing emotional stakes
- Robotic rhythm
- Adjective stacking

If no major issues are found, say the chapter is ready.

Create a clean Markdown change log with:

# Chapter {chapter_number} Audit Log

## Status
Ready / Needs Revision

## Issues Found
List issues clearly.

## Recommended Fixes
List fixes clearly.

## Notes
One short publishing-readiness note.

PROJECT FILES:

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}

CONFIRMED ENDING:
{ending_text}

OUTLINE:
{outline_text}

CHAPTER TEXT:
{chapter_text}
"""

def pick_recommended_option(options, recommended_number, label):
    """
    Picks the recommended option if found.
    Otherwise, returns the first option.
    """
    if recommended_number:
        for option in options:
            if re.search(rf"(?i)^{label}\s+{recommended_number}\b", option.strip()):
                return option

    return options[0] if options else ""


def generate_final_audit_prompt(metadata_text, character_text, ending_text, outline_text, manuscript_text):
    return f"""
You are a professional developmental editor and publishing-readiness auditor.

Perform a deep final audit of this novel.

Check for:
- Plot continuity
- Character consistency
- Timeline issues
- Contradictions
- Emotional arc problems
- Weak scenes
- Missing setup or payoff
- Names that conflict with the character profiles
- Publishing readiness

Do not rewrite the full manuscript.
Create a clear Markdown audit report with:
- Overall readiness score out of 100
- Critical issues
- Moderate issues
- Minor issues
- Specific chapter-level fixes needed
- Final publishing recommendation

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}

CONFIRMED ENDING:
{ending_text}

OUTLINE:
{outline_text}

MANUSCRIPT:
{manuscript_text}
"""

def generate_final_audit_and_fix_prompt(chapter_number, chapter_text, metadata_text, character_text, ending_text, outline_text, full_manuscript_text):
    return f"""
You are a professional developmental editor and final manuscript continuity fixer.

This is the final audit pass for the full novel.

Your task:
Audit Chapter {chapter_number} against the FULL MANUSCRIPT and project files.
If this chapter has any issue that affects publishing readiness, fix the chapter directly.

Your output must have exactly two sections:

# FIXED CHAPTER

[Write the complete corrected chapter here.]

# FINAL AUDIT CHANGE LOG

List every final-audit fix made to this chapter.

For each fix, include:
- Original problem
- Fix applied
- Why the fix was necessary

If no final-audit issues are found, keep the chapter unchanged and say:
No final-audit issues found. Chapter retained as-is.

Check especially for:
- Cross-chapter continuity errors
- Timeline problems
- Repeated or missing reveals
- Character motivation inconsistencies
- Relationship arc inconsistencies
- Chapter ending/beginning flow problems
- Setup/payoff issues
- Contradictions with the confirmed ending
- Incorrect names, ages, locations, relationship details
- Any remaining chapter title problems
- Any obvious publishing-readiness issue

PROJECT FILES:

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}

CONFIRMED ENDING:
{ending_text}

OUTLINE:
{outline_text}

FULL MANUSCRIPT:
{full_manuscript_text}

CHAPTER TO FIX:
{chapter_text}
"""

def final_audit_and_fix_all_chapters(metadata_text, character_text, ending_text, outline_text, model, status=None):
    """
    Runs a final audit/fix pass across all chapter files.
    Updates chapter files directly.
    Saves final audit fix logs.
    """

    chapter_files = sorted(CHAPTERS_DIR.glob("chapter_*.md"))

    chapter_files = [
        file for file in chapter_files
        if not file.name.endswith("_log.md") and not file.name.endswith("_final_log.md")
    ]

    full_manuscript_text = get_full_manuscript()
    final_logs = []

    for index, chapter_file in enumerate(chapter_files, start=1):
        chapter_match = re.search(r"chapter_(\d+)\.md", chapter_file.name, re.I)
        chapter_number = int(chapter_match.group(1)) if chapter_match else index

        if status:
            status.write(f"Step 7: Final auditing and fixing Chapter {chapter_number}...")

        chapter_text = chapter_file.read_text(encoding="utf-8")

        audit_fix_result = ask_openai(
            generate_final_audit_and_fix_prompt(
                chapter_number=chapter_number,
                chapter_text=chapter_text,
                metadata_text=metadata_text,
                character_text=character_text,
                ending_text=ending_text,
                outline_text=outline_text,
                full_manuscript_text=full_manuscript_text
            ),
            model=model
        )

        fixed_chapter, final_change_log = split_fixed_chapter_and_final_log(audit_fix_result)

        chapter_file.write_text(fixed_chapter, encoding="utf-8")

        final_log_file = CHAPTERS_DIR / f"chapter_{chapter_number:02}_final_log.md"
        final_log_file.write_text(final_change_log, encoding="utf-8")

        final_logs.append(f"# Chapter {chapter_number} Final Audit Log\n\n{final_change_log}")

        # Refresh manuscript after each chapter fix so later chapters compare against updated text
        full_manuscript_text = get_full_manuscript()

    combined_final_log = "\n\n".join(final_logs)
    save_markdown(OUTPUTS_DIR / "07_final_audit_fixes.md", combined_final_log)

    return combined_final_log


def generate_cover_prompt(metadata_text, character_text):
    return f"""
Create a book cover prompt for this novel.

Use this format:

Author Name: [VoidAndVelvet if MLM | AERESSA if WLW or others]

Please draw a book cover. It should have the title and the author name. Add enough space between text and the edge for print bleed.

TITLE:
[Insert title]

The story is about:
[Insert short description]

Book Cover Size: Portrait 6x9

PRIMARY Characters:
[Insert lead character 1: name, short description, physical description]
[Insert lead character 2: name, short description, physical description]
[Insert lead character 3 if applicable]
[Insert lead character 4 if applicable]

Additional Note:
The cover should be obviously LGBTQ+ and clearly match the story genre.

Art Style:
Semi-realistic digital illustration

Use the materials below.

METADATA:
{metadata_text}

CHARACTER PROFILES:
{character_text}
"""


def generate_draft2digital_prompt(manuscript_text, metadata_text):
    return f"""
Using the manuscript and metadata below, create complete SEO-optimized Draft2Digital publishing metadata.

Provide:

1. Short Description
- Single paragraph
- 50 to 400 characters
- No ending spoilers

2. Long Description
- Compelling book blurb
- Start with a hook
- Build tension and intrigue
- Introduce main characters and conflict
- Do not spoil the ending
- Do not mention other author names
- No trope lists, bullets, tags, keywords, content warnings, or extra metadata inside the long description

3. SEO Keywords
- 30 total
- Primary Keywords: 7
- Secondary Keywords: 13
- Long-Tail Keywords: 10

4. Maturity Level
Choose one:
- My book does NOT contain content inappropriate for minors.
- My book DOES contain content inappropriate for minors.
If considered erotica, say so.

5. BISAC Categories
Provide 5 total with full category codes and names.

6. Pricing Recommendation
Include ebook pricing suggestion, library pricing suggestion, and print pricing note.

Use clear Markdown headers.

METADATA:
{metadata_text}

MANUSCRIPT:
{manuscript_text}
"""


def generate_youtube_prompt(manuscript_text, metadata_text):
    return f"""
Act as a YouTube SEO specialist with expertise in audiobook marketing and LGBTQ+ romance content discovery.

Create comprehensive YouTube metadata optimized for discovery and click-through rate.

Input:
METADATA:
{metadata_text}

MANUSCRIPT:
{manuscript_text}

Create:

1. THREE TITLE VARIATIONS
Max 100 characters each.
- Version A: Search-optimized
- Version B: Trope-focused
- Version C: Full audiobook / audiobook to listen for sleep

The audiobook title must appear at the beginning if possible.

2. YOUTUBE DESCRIPTION

Description must start exactly with:

Get E-book Version: [LINK TO FOLLOW]
Buy Me a Coffee: https://buymeacoffee.com/aeressa

Then include:
- Short story overview, 3 to 5 short paragraphs, no spoilers
- Call to action
- Subscribe prompt
- Comment prompt asking about favorite moment
- Content warning in this format:
⚠️ Content includes: [warnings separated by commas]
- Hashtags, 10 to 15, at the very end

Use --- as horizontal separators between major description blocks.
Do not use --- between the two top links.

3. TWENTY SEO-OPTIMIZED TAGS
Comma-separated.
No hashtag symbols.
Include:
- 5 broad genre tags
- 5 specific trope tags
- 5 niche/subgenre tags
- 5 search behavior tags

4. PLAYLIST ASSIGNMENT
Choose from:
- MLM Romance Audiobook
- Sapphic Romance Audiobook
- Trans Romance Audiobooks
- Bisexual & Pansexual Romance Audiobooks
- Non-Binary Romance Audiobooks
- Polyamorous Queer Romance Audiobooks
- Asexual & Aromantic Romance Audiobooks
- Queer Romance Audiobooks (Master Playlist)

Indicate:
- PRIMARY playlist
- SECONDARY playlist(s), up to 2

Output as clean Markdown.
"""

def get_full_manuscript():
    """
    Combines all real chapter files into one manuscript string.
    Excludes audit/change log files.
    """
    chapter_files = sorted(CHAPTERS_DIR.glob("chapter_*.md"))

    chapter_files = [
        file for file in chapter_files
        if not file.name.endswith("_log.md")
    ]

    manuscript_parts = []

    for chapter_file in chapter_files:
        manuscript_parts.append(chapter_file.read_text(encoding="utf-8"))

    return "\n\n".join(manuscript_parts)

def clear_generated_chapters():
    """
    Deletes old generated chapter and chapter log files before a fresh Auto Mode run.
    Prevents old project chapters from being compiled into the new manuscript.
    """
    for file in CHAPTERS_DIR.glob("chapter_*.md"):
        file.unlink()

def extract_chapter_count_from_outline(outline_text):
    """
    Counts chapter headings from the Step 5 outline.
    Looks for formats like:
    Chapter 1
    # Chapter 1
    ## Chapter 1 - Title
    """
    matches = re.findall(
        r"(?im)^\s*#{0,6}\s*Chapter\s+\d+\b",
        outline_text
    )

    return len(matches)

# Streamlit UI
st.set_page_config(page_title="AI Novel Writing App", layout="wide")

st.title("AI Novel Writing App")
st.write("A simple app to automate your novel SOP step by step.")

st.sidebar.header("Settings")

model = st.sidebar.selectbox(
    "Model",
    [
        "gpt-5.5",
        "gpt-5.4",
        "gpt-5.4-mini"
    ],
    index=2
)

st.sidebar.caption("Use mini for cheaper drafts. Use stronger models for final chapter prose.")

page = st.sidebar.radio(
    "Choose Step",
    [
        "Step 1 - Pitch Maker",
        "Step 2 - Ending Plotting",
        "Step 3 - Metadata",
        "Step 4 - Character Profiles",
        "Step 5 - Outline",
        "Step 6 - Write Chapter",
        "Auto Mode - Step 1 to Step 10"
    ]
)

if page == "Step 1 - Pitch Maker":
    st.header("Step 1 - Pitch Maker")

    genre = st.text_input("Genre", value="sapphic dark romance / psychological thriller")
    story_idea = st.text_area("Paste your story idea here", height=300)

    if st.button("Generate Pitch Options"):
        if not story_idea.strip():
            st.error("Please paste a story idea first.")
        else:
            prompt = generate_pitch_prompt(story_idea, genre)
            with st.spinner("Generating pitch options..."):
                result = ask_openai(prompt, model=model)

            st.session_state["pitch_options_raw"] = result
            st.session_state["pitch_options"] = split_options(result, label="OPTION")

            recommended_pitch_number = extract_recommended_number(
                result,
                recommendation_label="OPTION",
                item_label="OPTION"
            )

            st.session_state["recommended_pitch_number"] = recommended_pitch_number

    if "pitch_options_raw" in st.session_state:
        st.subheader("Generated Pitch Options")
        st.markdown(st.session_state["pitch_options_raw"])

        pitch_options = st.session_state.get("pitch_options", [])

        if pitch_options:
            recommended_pitch_number = st.session_state.get("recommended_pitch_number")

            default_index = 0

            if recommended_pitch_number:
                for i, pitch in enumerate(pitch_options):
                    if re.search(rf"(?i)^OPTION\s+{recommended_pitch_number}\b", pitch.strip()):
                        default_index = i
                        break

            if recommended_pitch_number:
                st.info(f"Recommended pitch detected: Option {recommended_pitch_number}. Auto-selecting it by default.")

            selected_pitch = st.selectbox(
                "Select the pitch you want to use",
                pitch_options,
                index=default_index,
                format_func=lambda x: x.split("\n")[0][:80],
                key="selected_pitch"
            )

            if selected_pitch:
                save_markdown(OUTPUTS_DIR / "01_selected_pitch.md", selected_pitch)

                st.success("Selected pitch auto-saved to outputs/01_selected_pitch.md")
                st.text_area("Selected Pitch Preview", selected_pitch, height=300)
        else:
            st.warning("Could not split options cleanly. Copy your chosen pitch manually.")


if page == "Step 2 - Ending Plotting":
    st.header("Step 2 - Ending Plotting")

    pitch_text = read_markdown(OUTPUTS_DIR / "01_selected_pitch.md")

    if not pitch_text.strip():
        st.error("No selected pitch found. Go to Step 1 and select a pitch first.")
    else:
        st.subheader("Selected Pitch")
        st.text_area("Selected pitch file content", pitch_text, height=250)

        if st.button("Generate Ending Options"):
            prompt = generate_endings_prompt(pitch_text)
            with st.spinner("Generating ending options..."):
                result = ask_openai(prompt, model=model)

            st.session_state["ending_options_raw"] = result
            st.session_state["ending_options"] = split_options(result, label="ENDING")

            recommended_ending_number = extract_recommended_number(
                result,
                recommendation_label="ENDING",
                item_label="ENDING"
            )
            st.session_state["recommended_ending_number"] = recommended_ending_number

        if "ending_options_raw" in st.session_state:
            st.subheader("Generated Ending Options")
            st.markdown(st.session_state["ending_options_raw"])

            ending_options = st.session_state.get("ending_options", [])

            if ending_options:
                recommended_ending_number = st.session_state.get("recommended_ending_number")

                default_index = 0

                if recommended_ending_number:
                    for i, ending in enumerate(ending_options):
                        if re.search(rf"(?i)^ENDING\s+{recommended_ending_number}\b", ending.strip()):
                            default_index = i
                            break

                if recommended_ending_number:
                    st.info(f"Recommended ending detected: Ending {recommended_ending_number}. Auto-selecting it by default.")

                selected_ending = st.selectbox(
                    "Select the ending you want to use",
                    ending_options,
                    index=default_index,
                    format_func=lambda x: x.split("\n")[0][:80],
                    key="selected_ending"
                )

                st.text_area("Selected Ending Preview", selected_ending, height=300)

                selected_ending_hash = hashlib.md5(selected_ending.encode("utf-8")).hexdigest()

                if st.session_state.get("expanded_ending_hash") != selected_ending_hash:
                    with st.spinner("Expanding and auto-saving selected ending..."):
                        expand_prompt = generate_expand_ending_prompt(pitch_text, selected_ending)
                        expanded_ending = ask_openai(expand_prompt, model=model)

                    save_markdown(OUTPUTS_DIR / "02_selected_ending.md", expanded_ending)

                    st.session_state["expanded_ending_hash"] = selected_ending_hash
                    st.session_state["expanded_ending_text"] = expanded_ending

                expanded_ending_text = st.session_state.get(
                    "expanded_ending_text",
                    read_markdown(OUTPUTS_DIR / "02_selected_ending.md")
                )

                st.success("Expanded ending auto-saved to outputs/02_selected_ending.md")
                st.subheader("Expanded Confirmed Ending")
                st.markdown(expanded_ending_text)
            else:
                st.warning("Could not split endings cleanly. Copy your chosen ending manually.")


if page == "Step 3 - Metadata":
    st.header("Step 3 - Metadata")

    pitch_text = read_markdown(OUTPUTS_DIR / "01_selected_pitch.md")
    ending_text = read_markdown(OUTPUTS_DIR / "02_selected_ending.md")

    if st.button("Generate Metadata"):
        if not pitch_text.strip() or not ending_text.strip():
            st.error("Generate Step 1 and Step 2 first.")
        else:
            prompt = generate_metadata_prompt(pitch_text, ending_text)
            with st.spinner("Generating metadata..."):
                result = ask_openai(prompt, model=model)

            save_markdown(OUTPUTS_DIR / "03_metadata.md", result)
            st.success("Saved to outputs/03_metadata.md")
            st.markdown(result)


if page == "Step 4 - Character Profiles":
    st.header("Step 4 - Character Profiles")

    metadata_text = read_markdown(OUTPUTS_DIR / "03_metadata.md")
    ending_text = read_markdown(OUTPUTS_DIR / "02_selected_ending.md")

    if st.button("Generate Character Profiles"):
        if not metadata_text.strip() or not ending_text.strip():
            st.error("Generate metadata and endings first.")
        else:
            prompt = generate_character_prompt(metadata_text, ending_text)
            with st.spinner("Generating character profiles..."):
                result = ask_openai(prompt, model=model)

            save_markdown(OUTPUTS_DIR / "04_characters.md", result)
            st.success("Saved to outputs/04_characters.md")
            st.markdown(result)


if page == "Step 5 - Outline":
    st.header("Step 5 - Outline")

    metadata_text = read_markdown(OUTPUTS_DIR / "03_metadata.md")
    character_text = read_markdown(OUTPUTS_DIR / "04_characters.md")
    ending_text = read_markdown(OUTPUTS_DIR / "02_selected_ending.md")

    if st.button("Generate Outline"):
        if not metadata_text.strip() or not character_text.strip() or not ending_text.strip():
            st.error("Generate metadata, character profiles, and endings first.")
        else:
            prompt = generate_outline_prompt(metadata_text, character_text, ending_text)
            with st.spinner("Generating outline..."):
                result = ask_openai(prompt, model=model)

            save_markdown(OUTPUTS_DIR / "05_outline.md", result)
            st.success("Saved to outputs/05_outline.md")
            st.markdown(result)


if page == "Step 6 - Write Chapter":
    st.header("Step 6 - Write Chapter")

    chapter_number = st.number_input("Chapter number", min_value=1, step=1)

    metadata_text = read_markdown(OUTPUTS_DIR / "03_metadata.md")
    character_text = read_markdown(OUTPUTS_DIR / "04_characters.md")
    ending_text = read_markdown(OUTPUTS_DIR / "02_selected_ending.md")
    outline_text = read_markdown(OUTPUTS_DIR / "05_outline.md")

    if st.button("Write Chapter"):
        if not metadata_text.strip() or not character_text.strip() or not ending_text.strip() or not outline_text.strip():
            st.error("Generate metadata, characters, ending, and outline first.")
        else:
            prompt = generate_chapter_prompt(
                chapter_number,
                metadata_text,
                character_text,
                ending_text,
                outline_text
            )

            with st.spinner(f"Writing Chapter {chapter_number}..."):
                result = ask_openai(prompt, model=model)

            file_name = f"chapter_{int(chapter_number):02}.md"
            with st.spinner(f"Auditing and fixing Chapter {chapter_number}..."):
                audit_fix_result = ask_openai(
                    generate_chapter_audit_and_fix_prompt(
                        chapter_number,
                        result,
                        metadata_text,
                        character_text,
                        ending_text,
                        outline_text
                    ),
                    model=model
                )

            fixed_chapter, change_log = split_fixed_chapter_and_log(audit_fix_result)

            save_markdown(CHAPTERS_DIR / file_name, fixed_chapter)

            log_file_name = f"chapter_{int(chapter_number):02}_log.md"
            save_markdown(CHAPTERS_DIR / log_file_name, change_log)

            st.success(f"Fixed chapter saved to chapters/{file_name}")
            st.success(f"Change log saved to chapters/{log_file_name}")

            st.subheader("Fixed Chapter")
            st.markdown(fixed_chapter)

            st.subheader("Change Log")
            st.markdown(change_log)

if page == "Step 6 - Write Chapter":
    st.divider()
    st.subheader("Compile Manuscript")

    if st.button("Compile All Chapters to DOCX"):
        metadata_text = read_markdown(OUTPUTS_DIR / "03_metadata.md")
        selected_pitch_text = read_markdown(OUTPUTS_DIR / "01_selected_pitch.md")
        outline_text = read_markdown(OUTPUTS_DIR / "05_outline.md")

        if not metadata_text.strip():
            st.error("Metadata missing. Generate Step 3 first.")
        else:
            docx_path = compile_chapters_to_docx(
                metadata_text=metadata_text,
                selected_pitch_text=selected_pitch_text,
                outline_text=outline_text
            )
            st.success(f"Compiled manuscript saved to {docx_path}")

if page == "Auto Mode - Step 1 to Step 10":
    st.header("Auto Mode - Step 1 to Step 10")

    st.warning(
        "Auto Mode will call the API multiple times. This is convenient, but it can use more credits."
    )

    auto_enabled = st.toggle("Enable Auto Mode")

    genre = st.text_input(
        "Genre",
        value="sapphic dark romance / psychological thriller",
        key="auto_genre"
    )

    story_idea = st.text_area(
        "Paste your story idea here",
        height=300,
        key="auto_story_idea"
    )

    existing_outline_text = read_markdown(OUTPUTS_DIR / "05_outline.md")
    existing_outline_chapter_count = extract_chapter_count_from_outline(existing_outline_text)

    suggested_chapter_count = max(20, existing_outline_chapter_count)

    st.info(
        f"Auto Mode will use the chapter count from Step 5 outline. "
        f"Current detected count: {existing_outline_chapter_count if existing_outline_chapter_count else 'No outline found yet'}. "
        f"Minimum chapter count: 20."
    )

    fallback_chapter_count = st.number_input(
        "Fallback chapter count if Auto Mode cannot detect chapters from the outline",
        min_value=20,
        max_value=100,
        value=suggested_chapter_count,
        step=1
    )

    final_output_parent_folder = st.text_input(
        "Final publishing output folder",
        value=str((BASE_DIR / "final_exports").resolve()),
        help="Paste the folder path where you want the final publishing files saved."
    )

    run_auto = st.button("Run Auto Mode from Step 1 to Step 10")

    if run_auto:
        if not auto_enabled:
            st.error("Turn on the Auto Mode toggle first.")
        elif not story_idea.strip():
            st.error("Please paste your story idea first.")
        else:
            progress = st.progress(0)
            status = st.empty()

            clear_generated_chapters()

            # STEP 1 - Pitch Maker
            status.write("Step 1: Generating pitch options...")
            pitch_options_raw = ask_openai(
                generate_pitch_prompt(story_idea, genre),
                model=model
            )

            pitch_options = split_options(pitch_options_raw, label="OPTION")

            recommended_pitch_number = extract_recommended_number(
                pitch_options_raw,
                recommendation_label="OPTION",
                item_label="OPTION"
            )

            selected_pitch = pick_recommended_option(
                pitch_options,
                recommended_pitch_number,
                label="OPTION"
            )

            if not selected_pitch.strip():
                st.error("Auto Mode stopped: could not detect a selected pitch from Step 1.")
                st.stop()

            save_markdown(OUTPUTS_DIR / "01_pitch_options_raw.md", pitch_options_raw)
            save_markdown(OUTPUTS_DIR / "01_selected_pitch.md", selected_pitch)

            progress.progress(10)

            # STEP 2 - Ending Plotting
            status.write("Step 2: Generating ending options...")
            ending_options_raw = ask_openai(
                generate_endings_prompt(selected_pitch),
                model=model
            )

            ending_options = split_options(ending_options_raw, label="ENDING")

            recommended_ending_number = extract_recommended_number(
                ending_options_raw,
                recommendation_label="ENDING",
                item_label="ENDING"
            )

            selected_ending = pick_recommended_option(
                ending_options,
                recommended_ending_number,
                label="ENDING"
            )

            if not selected_ending.strip():
                st.error("Auto Mode stopped: could not detect a selected ending from Step 2.")
                st.stop()

            status.write("Step 2: Expanding selected ending...")
            expanded_ending = ask_openai(
                generate_expand_ending_prompt(selected_pitch, selected_ending),
                model=model
            )

            save_markdown(OUTPUTS_DIR / "02_ending_options_raw.md", ending_options_raw)
            save_markdown(OUTPUTS_DIR / "02_selected_ending.md", expanded_ending)

            progress.progress(20)

            # STEP 3 - Metadata
            status.write("Step 3: Generating metadata...")
            metadata = ask_openai(
                generate_metadata_prompt(selected_pitch, expanded_ending),
                model=model
            )
            save_markdown(OUTPUTS_DIR / "03_metadata.md", metadata)

            progress.progress(30)

            # STEP 4 - Character Profiles
            status.write("Step 4: Generating character profiles...")
            characters = ask_openai(
                generate_character_prompt(metadata, expanded_ending),
                model=model
            )
            save_markdown(OUTPUTS_DIR / "04_characters.md", characters)

            progress.progress(40)

            # STEP 5 - Outline
            status.write("Step 5: Generating outline...")
            outline = ask_openai(
                generate_outline_prompt(metadata, characters, expanded_ending),
                model=model
            )
            save_markdown(OUTPUTS_DIR / "05_outline.md", outline)

            detected_chapter_count = extract_chapter_count_from_outline(outline)
            chapter_count = max(20, detected_chapter_count or int(fallback_chapter_count))

            status.write(f"Step 5 complete. Detected {detected_chapter_count} chapters from outline. Auto Mode will write {chapter_count} chapters.")

            progress.progress(50)

            # STEP 6 - Write Chapters
            status.write("Step 6: Writing chapters...")

            for chapter_number in range(1, int(chapter_count) + 1):
                status.write(f"Step 6: Writing Chapter {chapter_number} of {chapter_count}...")

                chapter = ask_openai(
                    generate_chapter_prompt(
                        chapter_number,
                        metadata,
                        characters,
                        expanded_ending,
                        outline
                    ),
                    model=model
                )

                status.write(f"Step 6: Auditing and fixing Chapter {chapter_number}...")

                audit_fix_result = ask_openai(
                    generate_chapter_audit_and_fix_prompt(
                        chapter_number,
                        chapter,
                        metadata,
                        characters,
                        expanded_ending,
                        outline
                    ),
                    model=model
                )

                fixed_chapter, change_log = split_fixed_chapter_and_log(audit_fix_result)

                file_name = f"chapter_{chapter_number:02}.md"
                save_markdown(CHAPTERS_DIR / file_name, fixed_chapter)

                log_file_name = f"chapter_{chapter_number:02}_log.md"
                save_markdown(CHAPTERS_DIR / log_file_name, change_log)


                chapter_progress = 50 + int((chapter_number / int(chapter_count)) * 20)
                progress.progress(min(chapter_progress, 70))

            manuscript = get_full_manuscript()

            # STEP 7 - Final Audit + Auto-Fix
            status.write("Step 7: Running final audit report...")

            manuscript = get_full_manuscript()

            final_audit = ask_openai(
                generate_final_audit_prompt(
                    metadata,
                    characters,
                    expanded_ending,
                    outline,
                    manuscript
                ),
                model=model
            )

            save_markdown(OUTPUTS_DIR / "07_final_audit.md", final_audit)

            status.write("Step 7: Auto-fixing final manuscript issues chapter by chapter...")

            final_audit_fixes = final_audit_and_fix_all_chapters(
                metadata_text=metadata,
                character_text=characters,
                ending_text=expanded_ending,
                outline_text=outline,
                model=model,
                status=status
            )

            save_markdown(OUTPUTS_DIR / "07_final_audit_fixes.md", final_audit_fixes)

            # Refresh manuscript after final fixes
            manuscript = get_full_manuscript()

            progress.progress(80)

            status.write("Compiling internal final manuscript DOCX...")
            selected_pitch_text = selected_pitch

            docx_path = compile_chapters_to_docx(
                metadata_text=metadata,
                selected_pitch_text=selected_pitch_text,
                outline_text=outline
            )

            save_markdown(PUBLISHING_DIR / "final_manuscript_path.txt", str(docx_path))

            # STEP 8 - Book Cover Prompt
            status.write("Step 8: Generating book cover prompt...")
            cover_prompt = ask_openai(
                generate_cover_prompt(metadata, characters),
                model=model
            )
            save_markdown(PUBLISHING_DIR / "08_book_cover_prompt.md", cover_prompt)

            progress.progress(86)

            # STEP 9 - Draft2Digital Details
            status.write("Step 9: Generating Draft2Digital metadata...")
            d2d_metadata = ask_openai(
                generate_draft2digital_prompt(manuscript, metadata),
                model=model
            )
            save_markdown(PUBLISHING_DIR / "09_draft2digital_metadata.md", d2d_metadata)

            progress.progress(93)

            # STEP 10 - YouTube Details
            status.write("Step 10: Generating YouTube metadata...")
            youtube_metadata = ask_openai(
                generate_youtube_prompt(manuscript, metadata),
                model=model
            )
            save_markdown(PUBLISHING_DIR / "10_youtube_metadata.md", youtube_metadata)

            status.write("Exporting final publishing files...")

            exported_files = export_final_publishing_outputs(
                parent_folder=final_output_parent_folder,
                metadata_text=metadata,
                selected_pitch_text=selected_pitch,
                outline_text=outline,
                cover_prompt=cover_prompt,
                d2d_metadata=d2d_metadata,
                youtube_metadata=youtube_metadata
            )

            save_markdown(PUBLISHING_DIR / "final_export_folder.txt", str(exported_files["folder"]))

            progress.progress(100)

            status.write("Auto Mode complete.")

            st.success("Done. Auto Mode finished Step 1 to Step 10.")

            st.subheader("Saved Files")
            st.markdown("""
                - `outputs/01_pitch_options_raw.md`
                - `outputs/01_selected_pitch.md`
                - `outputs/02_ending_options_raw.md`
                - `outputs/02_selected_ending.md`
                - `outputs/03_metadata.md`
                - `outputs/04_characters.md`
                - `outputs/05_outline.md`
                - `outputs/07_final_audit.md`
                - `chapters/chapter_01.md` and onward
                - `publishing/08_book_cover_prompt.md`
                - `publishing/09_draft2digital_metadata.md`
                - `publishing/10_youtube_metadata.md`
                - `publishing/final_manuscript.docx`
                - `outputs/07_final_audit_fixes.md`
                - `chapters/chapter_01_final_log.md` and onward
                """)
            
            st.subheader("Final Publishing Export Folder")
            st.write(str(exported_files["folder"]))

            st.markdown(f"""
                - Final Manuscript: `{exported_files["manuscript"]}`
                - Book Cover Prompt: `{exported_files["cover_prompt"]}`
                - Draft2Digital Details: `{exported_files["draft2digital"]}`
                - YouTube Details: `{exported_files["youtube"]}`
                """)

            st.subheader("Selected Pitch")
            st.markdown(selected_pitch)

            st.subheader("Confirmed Ending")
            st.markdown(expanded_ending)

            st.subheader("Final Audit")
            st.markdown(final_audit)